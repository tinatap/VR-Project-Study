import socket
import struct
import json
import os
import threading
import time
from datetime import datetime

from openpyxl import Workbook, load_workbook


# ============================================================
# SETTINGS
# ============================================================

HOST = "0.0.0.0"
PORT = 12345

DATA_FOLDER = "analytics_data"

os.makedirs(DATA_FOLDER, exist_ok=True)


# ============================================================
# CONTINUOUS DATA FILES
# ============================================================

DATA_FILE = os.path.join(
    DATA_FOLDER,
    "continuous_data.jsonl"
)

EXCEL_FILE = os.path.join(
    DATA_FOLDER,
    "continuous_data.xlsx"
)


# ============================================================
# EXPERIMENT SUMMARY FILE
# ============================================================

SUMMARY_FILE = os.path.join(
    DATA_FOLDER,
    "experiment_summary.json"
)


# ============================================================
# EXCEL SETTINGS
# ============================================================

EXCEL_SAVE_INTERVAL = 10


# ============================================================
# EXCEL VARIABLES
# ============================================================

excel_lock = threading.Lock()

excel_dirty = False

excel_running = True


# ============================================================
# CREATE EXCEL
# ============================================================

print()
print("Creating Excel file...")

try:

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Data"

    workbook.save(
        EXCEL_FILE
    )

    workbook.close()

    print(
        "Excel file created successfully:"
    )

    print(
        EXCEL_FILE
    )

except Exception as e:

    print(
        "!!! EXCEL CREATION ERROR !!!"
    )

    print(
        repr(e)
    )

    raise


# ============================================================
# OPEN EXCEL
# ============================================================

workbook = load_workbook(
    EXCEL_FILE
)

sheet = workbook["Data"]

headers = {}


# ============================================================
# ADD HEADER
# ============================================================

def get_column(key):

    global headers

    if key in headers:
        return headers[key]


    # --------------------------------------------------------
    # New column
    # --------------------------------------------------------

    column = sheet.max_column + 1


    # Empty first cell

    if (
        sheet.max_row == 1
        and sheet.cell(
            row=1,
            column=1
        ).value is None
    ):

        column = 1


    sheet.cell(
        row=1,
        column=column
    ).value = key

    headers[key] = column

    return column


# ============================================================
# ADD DATA TO EXCEL
# ============================================================

def add_to_excel(data):

    global excel_dirty

    with excel_lock:

        try:

            # ------------------------------------------------
            # Create columns
            # ------------------------------------------------

            for key in data.keys():

                get_column(key)


            # ------------------------------------------------
            # New row
            # ------------------------------------------------

            row = sheet.max_row + 1


            # ------------------------------------------------
            # Write values
            # ------------------------------------------------

            for key, value in data.items():

                column = headers[key]

                cell = sheet.cell(
                    row=row,
                    column=column
                )


                # --------------------------------------------
                # Boolean
                # --------------------------------------------

                if isinstance(value, bool):

                    cell.value = value


                # --------------------------------------------
                # Number
                # --------------------------------------------

                elif isinstance(
                    value,
                    (int, float)
                ):

                    cell.value = value


                # --------------------------------------------
                # None
                # --------------------------------------------

                elif value is None:

                    cell.value = None


                # --------------------------------------------
                # List / Dictionary
                # --------------------------------------------

                elif isinstance(
                    value,
                    (list, dict)
                ):

                    cell.value = json.dumps(
                        value,
                        ensure_ascii=False
                    )


                # --------------------------------------------
                # String
                # --------------------------------------------

                else:

                    cell.value = value


            excel_dirty = True


        except Exception as e:

            print()
            print(
                "!!! EXCEL DATA ERROR !!!"
            )

            print(
                repr(e)
            )


# ============================================================
# EXCEL AUTO SAVE
# ============================================================

def excel_auto_save():

    global excel_dirty

    print(
        f"Excel auto-save started: "
        f"every {EXCEL_SAVE_INTERVAL} seconds"
    )


    while excel_running:

        time.sleep(
            EXCEL_SAVE_INTERVAL
        )


        with excel_lock:

            if excel_dirty:

                try:

                    workbook.save(
                        EXCEL_FILE
                    )

                    excel_dirty = False

                    print()

                    print(
                        "[EXCEL] Saved successfully."
                    )


                except Exception as e:

                    print()

                    print(
                        "!!! EXCEL SAVE ERROR !!!"
                    )

                    print(
                        repr(e)
                    )


# ============================================================
# START EXCEL THREAD
# ============================================================

excel_thread = threading.Thread(
    target=excel_auto_save,
    daemon=True
)

excel_thread.start()


# ============================================================
# RECEIVE EXACT BYTES
# ============================================================

def receive_exact(
    conn,
    size
):

    data = b""


    while len(data) < size:

        chunk = conn.recv(
            size - len(data)
        )


        if not chunk:

            return None


        data += chunk


    return data


# ============================================================
# RECEIVE MESSAGE
# ============================================================

def receive_message(conn):

    # --------------------------------------------------------
    # First 4 bytes = message length
    # --------------------------------------------------------

    length_bytes = receive_exact(
        conn,
        4
    )


    if length_bytes is None:

        return None


    # --------------------------------------------------------
    # Big Endian unsigned integer
    # Compatible with Unity:
    #
    # IPAddress.HostToNetworkOrder(...)
    #
    # --------------------------------------------------------

    message_length = struct.unpack(
        "!I",
        length_bytes
    )[0]


    if message_length <= 0:

        print(
            "Invalid message length:",
            message_length
        )

        return None


    if message_length > 10 * 1024 * 1024:

        print(
            "Message too large:",
            message_length
        )

        return None


    # --------------------------------------------------------
    # Receive JSON bytes
    # --------------------------------------------------------

    json_bytes = receive_exact(
        conn,
        message_length
    )


    if json_bytes is None:

        return None


    try:

        json_string = json_bytes.decode(
            "utf-8"
        )


        data = json.loads(
            json_string
        )


        return data


    except Exception as e:

        print(
            "JSON ERROR:",
            repr(e)
        )

        return None


# ============================================================
# SAVE CONTINUOUS JSON
# ============================================================

def save_continuous_json(data):

    with open(
        DATA_FILE,
        "a",
        encoding="utf-8"
    ) as file:

        file.write(
            json.dumps(
                data,
                ensure_ascii=False,
                separators=(",", ":")
            )
        )

        file.write("\n")


# ============================================================
# SAVE EXPERIMENT SUMMARY
# ============================================================

def save_experiment_summary(data):

    try:

        # ----------------------------------------------------
        # Add server time
        # ----------------------------------------------------

        data["serverReceivedTime"] = (
            datetime.now().isoformat()
        )


        # ----------------------------------------------------
        # Save as formatted JSON
        # ----------------------------------------------------

        with open(
            SUMMARY_FILE,
            "w",
            encoding="utf-8"
        ) as file:

            json.dump(
                data,
                file,
                ensure_ascii=False,
                indent=4
            )


        # ----------------------------------------------------
        # Console information
        # ----------------------------------------------------

        maze_visits = data.get(
            "mazeVisits",
            []
        )


        print()
        print(
            "=========================================="
        )

        print(
            "EXPERIMENT SUMMARY RECEIVED"
        )

        print(
            "=========================================="
        )

        print(
            "Final Result:",
            data.get("finalResult")
        )

        print(
            "Total Game Time:",
            data.get("totalGameTime")
        )

        print(
            "Final Score:",
            data.get("finalScore")
        )

        print(
            "Highest Completed Maze:",
            data.get("highestCompletedMaze")
        )

        print(
            "Start Room Duration:",
            data.get("startRoomDuration")
        )

        print(
            "Start Question Panel Duration:",
            data.get("startQuestionPanelDuration")
        )

        print(
            "Maze Visits:",
            len(maze_visits)
        )

        print(
            "Summary file:",
            SUMMARY_FILE
        )

        print(
            "=========================================="
        )


        # ----------------------------------------------------
        # Print each maze visit
        # ----------------------------------------------------

        for visit in maze_visits:

            print(
                "Visit #{visit} | "
                "Maze={maze} | "
                "Attempt={attempt} | "
                "Coins={coins}/{totalCoins} | "
                "Duration={duration:.3f}s | "
                "Result={result} | "
                "End={end}".format(

                    visit=visit.get(
                        "visitNumber"
                    ),

                    maze=visit.get(
                        "mazeNumber"
                    ),

                    attempt=visit.get(
                        "attemptNumber"
                    ),

                    coins=visit.get(
                        "collectedCoins",
                        0
                    ),

                    totalCoins=visit.get(
                        "totalCoins",
                        0
                    ),

                    duration=float(
                        visit.get(
                            "durationSeconds",
                            0
                        )
                    ),

                    result=visit.get(
                        "result"
                    ),

                    end=visit.get(
                        "endReason"
                    )
                )
            )


        print(
            "=========================================="
        )


    except Exception as e:

        print()
        print(
            "!!! SUMMARY SAVE ERROR !!!"
        )

        print(
            repr(e)
        )


# ============================================================
# HANDLE CONTINUOUS DATA
# ============================================================

def handle_continuous_data(data):

    # --------------------------------------------------------
    # Add server reception time
    # --------------------------------------------------------

    data["serverReceivedTime"] = (
        datetime.now().isoformat()
    )


    # --------------------------------------------------------
    # Save JSONL
    # --------------------------------------------------------

    save_continuous_json(
        data
    )


    # --------------------------------------------------------
    # Save Excel
    # --------------------------------------------------------

    add_to_excel(
        data
    )


# ============================================================
# HANDLE CLIENT
# ============================================================

def handle_client(
    conn,
    address
):

    print()

    print(
        "=========================================="
    )

    print(
        "Unity connected!"
    )

    print(
        "Client:",
        address
    )

    print(
        "=========================================="
    )


    message_count = 0


    try:

        while True:

            data = receive_message(
                conn
            )


            if data is None:

                print(
                    "Unity disconnected."
                )

                break


            message_count += 1


            # =================================================
            # IDENTIFY MESSAGE TYPE
            # =================================================

            record_type = data.get(
                "recordType",
                ""
            )


            message_type = data.get(
                "messageType",
                ""
            )


            # =================================================
            # EXPERIMENT SUMMARY
            # =================================================

            if (
                record_type ==
                "EXPERIMENT_SUMMARY"
                or
                message_type ==
                "EXPERIMENT_SUMMARY"
            ):

                save_experiment_summary(
                    data
                )

                continue


            # =================================================
            # CONTINUOUS DATA / EVENT
            # =================================================

            handle_continuous_data(
                data
            )


            # =================================================
            # CONSOLE
            # =================================================

            print(
                f"[{message_count}] "
                f"Type={record_type} | "
                f"Maze={data.get('mazeNumber')} | "
                f"Attempt={data.get('attemptNumber')}"
            )


    except ConnectionResetError:

        print(
            "Connection reset by Unity."
        )


    except Exception as e:

        print(
            "CLIENT ERROR:"
        )

        print(
            repr(e)
        )


    finally:

        # ----------------------------------------------------
        # FINAL EXCEL SAVE
        # ----------------------------------------------------

        with excel_lock:

            try:

                workbook.save(
                    EXCEL_FILE
                )

                print(
                    "Excel final save completed."
                )


            except Exception as e:

                print(
                    "!!! FINAL EXCEL SAVE ERROR !!!"
                )

                print(
                    repr(e)
                )


        try:

            conn.close()

        except:

            pass


        print(
            "Connection closed:",
            address
        )


# ============================================================
# START SERVER
# ============================================================

def start_server():

    global excel_running


    print()

    print(
        "=========================================="
    )

    print(
        "PYTHON TCP SERVER"
    )

    print(
        "=========================================="
    )

    print(
        f"Listening on {HOST}:{PORT}"
    )

    print()

    print(
        "CONTINUOUS JSON:",
        DATA_FILE
    )

    print(
        "CONTINUOUS EXCEL:",
        EXCEL_FILE
    )

    print()

    print(
        "EXPERIMENT SUMMARY:",
        SUMMARY_FILE
    )

    print()

    print(
        "Excel save interval:",
        EXCEL_SAVE_INTERVAL,
        "seconds"
    )

    print(
        "=========================================="
    )


    # ========================================================
    # CREATE SERVER
    # ========================================================

    server = socket.socket(
        socket.AF_INET,
        socket.SOCK_STREAM
    )


    server.setsockopt(
        socket.SOL_SOCKET,
        socket.SO_REUSEADDR,
        1
    )


    server.bind(
        (
            HOST,
            PORT
        )
    )


    server.listen(
        1
    )


    print(
        "Waiting for Unity..."
    )


    try:

        while True:

            conn, address = server.accept()


            handle_client(
                conn,
                address
            )


            print()

            print(
                "Waiting for Unity again..."
            )


    except KeyboardInterrupt:

        print(
            "\nServer stopped."
        )


    finally:

        print(
            "Final Excel save..."
        )


        excel_running = False


        with excel_lock:

            try:

                workbook.save(
                    EXCEL_FILE
                )

                print(
                    "Final Excel save completed."
                )


            except Exception as e:

                print(
                    "!!! FINAL SAVE ERROR !!!"
                )

                print(
                    repr(e)
                )


            try:

                workbook.close()

            except:

                pass


        try:

            server.close()

        except:

            pass


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":

    start_server()